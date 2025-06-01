import os
from pathlib import Path
from docx import Document
from openpyxl import Workbook

def extract_students_from_docx(docx_path):
    doc = Document(docx_path)
    students = []
    group_name = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if text.upper().startswith("ГРУППА"):
            parts = text.strip().split()
            if len(parts) >= 2:
                group_name = parts[1].strip().upper()
            continue

        fio_parts = text.split()
        if len(fio_parts) < 2:
            continue
        if len(fio_parts) == 2:
            fio_parts.append("")

        last_name = fio_parts[0].capitalize()
        first_name = fio_parts[1].capitalize()
        middle_name = fio_parts[2].capitalize() if fio_parts[2] else ""

        students.append([
            first_name,
            middle_name,
            last_name,
            "2005-09-01",
            "ЛОЖЬ",
            group_name or "",
            "2022-09-01"
        ])
    return students, group_name or "UNKNOWN"

def save_students_to_xlsx(students, group_name, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Студенты"
    headers = ["first_name", "middle_name", "last_name", "date_birthday", "is_headman", "group", "date_input"]
    ws.append(headers)
    for row in students:
        ws.append(row)
    wb.save(output_path)

def process_all_docx_in_folder(folder_path):
    folder = Path(folder_path)
    docx_files = list(folder.glob("*.docx"))
    print(f"🔍 Найдено файлов: {len(docx_files)}")

    for docx_file in docx_files:
        students, group_name = extract_students_from_docx(docx_file)
        if not students:
            print(f"⚠️ Пропущен (нет студентов): {docx_file.name}")
            continue

        output_path = docx_file.with_suffix(".xlsx")
        save_students_to_xlsx(students, group_name, output_path)
        print(f"✅ Обработан: {docx_file.name} → {output_path.name}")

# Пример запуска

