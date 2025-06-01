from docx import Document
from openpyxl import Workbook

def parse_docx_to_student_xlsx(docx_file, output_path):
    document = Document(docx_file)
    students = []
    group_name = None

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if text.upper().startswith("ГРУППА"):
            parts = text.split()
            if len(parts) >= 2:
                group_name = parts[1].strip().upper()
            continue

        fio_parts = text.split()
        if len(fio_parts) < 2:
            continue
        if len(fio_parts) == 2:
            fio_parts.append("")

        students.append([
            fio_parts[1].capitalize(),
            fio_parts[2].capitalize() if len(fio_parts) > 2 else "",
            fio_parts[0].capitalize(),
            "2005-09-01",
            "ЛОЖЬ",
            group_name or "",
            "2022-09-01"
        ])

    wb = Workbook()
    ws = wb.active
    ws.append(["first_name", "middle_name", "last_name", "date_birthday", "is_headman", "group", "date_input"])
    for student in students:
        ws.append(student)
    wb.save(output_path)





import random
import string
from django.contrib.auth import get_user_model
from users.models import DisciplinePrepod, TutorPrepod

User = get_user_model()


def generate_password(length=8):
    """Генерация случайного пароля из букв и цифр."""
    characters = string.ascii_letters + string.digits
    return ''.join(random.choice(characters) for _ in range(length))


def generate_prepod_login():
    """Генерация уникального логина с префиксом prepods_"""
    while True:
        username = "prepods_" + ''.join(random.choices(string.ascii_lowercase + string.digits, k=6))
        if not User.objects.filter(email=username + "@example.com").exists():
            return username + "@example.com"


def create_prepod_user(email, password):
    """Создание пользователя-преподавателя с указанным email и паролем."""
    user = User.objects.create_user(email=email, password=password)
    user.is_active = True
    user.save()
    return user


def create_discipline_prepod(user):
    """Создание экземпляра DisciplinePrepod для пользователя."""
    return DisciplinePrepod.objects.create(user=user)


def create_tutor_prepod(user):
    """Создание экземпляра TutorPrepod для пользователя."""
    return TutorPrepod.objects.create(user=user)
